"""Génération des fichiers Word, Excel et PDF."""
import io
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table,
                                TableStyle, Image, PageBreak)

from services.meteo import MOIS_NOMS, VARIABLES


MOIS_COURTS = ["Jan", "Fév", "Mar", "Avr", "Mai", "Juin",
               "Juil", "Août", "Sep", "Oct", "Nov", "Déc"]


# ============================================================
# Graphiques
# ============================================================
def generer_graphique_ombrothermique(pivots: dict) -> io.BytesIO:
    """Diagramme classique température + pluie (si les deux sont présents)."""
    if "temperature" not in pivots or "pluie" not in pivots:
        return None

    fig, ax1 = plt.subplots(figsize=(10, 5))
    ax1.bar(MOIS_COURTS, pivots["pluie"]["Moyenne"].values,
            color="#1f77b4", alpha=0.6, label="Précipitations")
    ax1.set_xlabel("Mois")
    ax1.set_ylabel("Précipitations (mm)", color="#1f77b4")
    ax1.tick_params(axis="y", labelcolor="#1f77b4")

    ax2 = ax1.twinx()
    ax2.plot(MOIS_COURTS, pivots["temperature"]["Moyenne"].values,
             color="#d62728", marker="o", linewidth=2, label="Température")
    ax2.set_ylabel("Température (°C)", color="#d62728")
    ax2.tick_params(axis="y", labelcolor="#d62728")

    plt.title("Diagramme ombrothermique — Moyennes")
    fig.tight_layout()

    buffer = io.BytesIO()
    plt.savefig(buffer, format="png", dpi=150, bbox_inches="tight")
    plt.close()
    buffer.seek(0)
    return buffer


def generer_graphique_simple(pivot: pd.DataFrame, variable: str) -> io.BytesIO:
    """Graphique mensuel simple pour une variable secondaire."""
    config = VARIABLES[variable]
    fig, ax = plt.subplots(figsize=(10, 4))

    couleurs = {
        "humidite": "#17becf",
        "vent": "#7f7f7f",
        "ensoleillement": "#ff7f0e",
    }
    couleur = couleurs.get(variable, "#2ca02c")

    ax.bar(MOIS_COURTS, pivot["Moyenne"].values, color=couleur, alpha=0.75)
    ax.set_xlabel("Mois")
    ax.set_ylabel(config["label"])
    ax.set_title(f"{config['label']} — Moyenne mensuelle")
    ax.grid(axis="y", linestyle="--", alpha=0.4)
    fig.tight_layout()

    buffer = io.BytesIO()
    plt.savefig(buffer, format="png", dpi=150, bbox_inches="tight")
    plt.close()
    buffer.seek(0)
    return buffer


def generer_graphique_comparaison(pivot: pd.DataFrame, variable: str) -> io.BytesIO:
    """Courbes superposées : une courbe par lieu."""
    config = VARIABLES[variable]
    fig, ax = plt.subplots(figsize=(11, 5))

    palette = ["#1f77b4", "#d62728", "#2ca02c", "#ff7f0e", "#9467bd"]

    for i, lieu in enumerate(pivot.columns):
        ax.plot(MOIS_COURTS, pivot[lieu].values,
                color=palette[i % len(palette)],
                marker="o", linewidth=2, label=lieu)

    ax.set_xlabel("Mois")
    ax.set_ylabel(config["label"])
    ax.set_title(f"Comparaison — {config['label']}")
    ax.legend(loc="best", framealpha=0.9)
    ax.grid(linestyle="--", alpha=0.4)
    fig.tight_layout()

    buffer = io.BytesIO()
    plt.savefig(buffer, format="png", dpi=150, bbox_inches="tight")
    plt.close()
    buffer.seek(0)
    return buffer


# ============================================================
# Helpers
# ============================================================
def _pivot_vers_lignes(pivot: pd.DataFrame):
    entetes = ["Mois"] + [str(c) for c in pivot.columns]
    lignes = []
    for mois_num in pivot.index:
        ligne = [MOIS_NOMS[mois_num - 1]]
        for col in pivot.columns:
            ligne.append(f"{pivot.loc[mois_num, col]:.1f}")
        lignes.append(ligne)
    return entetes, lignes


def _ajouter_tableau_word(doc, pivot):
    entetes, lignes = _pivot_vers_lignes(pivot)
    table = doc.add_table(rows=1, cols=len(entetes))
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for i, txt in enumerate(entetes):
        cell = table.rows[0].cells[i]
        cell.text = txt
        for para in cell.paragraphs:
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in para.runs:
                run.bold = True

    for ligne in lignes:
        cells = table.add_row().cells
        for i, val in enumerate(ligne):
            cells[i].text = val
            if i > 0:
                for para in cells[i].paragraphs:
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _ajouter_extremes_word(doc, extremes, titre_section="Statistiques d'extrêmes"):
    if not extremes:
        return
    doc.add_heading(titre_section, level=2)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    entetes = ["Indicateur", "Valeur", "Date / Période"]
    for i, txt in enumerate(entetes):
        cell = table.rows[0].cells[i]
        cell.text = txt
        for para in cell.paragraphs:
            for run in para.runs:
                run.bold = True
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for cle, e in extremes.items():
        row = table.add_row().cells
        row[0].text = f"{e['icone']} {e['label']}"
        row[1].text = e["valeur"]
        row[2].text = e["date"]
        for para in row[1].paragraphs:
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _ajouter_extremes_pdf(elements, extremes, style_h2,
                            titre_section="Statistiques d'extrêmes"):
    if not extremes:
        return
    elements.append(Paragraph(titre_section, style_h2))
    data = [["Indicateur", "Valeur", "Date / Période"]]
    for cle, e in extremes.items():
        data.append([f"{e['icone']} {e['label']}", e["valeur"], e["date"]])
    table = Table(data, repeatRows=1, colWidths=[8*cm, 5*cm, 9*cm])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4472C4")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("ALIGN", (1, 0), (1, -1), "CENTER"),
        ("FONTSIZE", (0, 0), (-1, -1), 10),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1),
         [colors.white, colors.HexColor("#F2F2F2")]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))
    elements.append(table)
    elements.append(PageBreak())


def _ajouter_extremes_excel(wb, extremes, nom_feuille="Extrêmes"):
    if not extremes:
        return
    ws = wb.create_sheet(nom_feuille[:31])
    ws["A1"] = nom_feuille
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells("A1:C1")

    entetes = ["Indicateur", "Valeur", "Date / Période"]
    fill = PatternFill("solid", fgColor="D9E1F2")
    for col_idx, txt in enumerate(entetes, start=1):
        cell = ws.cell(row=3, column=col_idx, value=txt)
        cell.font = Font(bold=True)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center")

    for row_idx, (cle, e) in enumerate(extremes.items(), start=4):
        ws.cell(row=row_idx, column=1, value=f"{e['icone']} {e['label']}")
        ws.cell(row=row_idx, column=2, value=e["valeur"])
        ws.cell(row=row_idx, column=3, value=e["date"])
        ws.cell(row=row_idx, column=2).alignment = Alignment(horizontal="center")

    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 30


def _remplir_feuille_excel(ws, titre, pivot):
    ws["A1"] = titre
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells(start_row=1, start_column=1,
                   end_row=1, end_column=len(pivot.columns) + 1)

    entetes, lignes = _pivot_vers_lignes(pivot)
    fill = PatternFill("solid", fgColor="D9E1F2")
    centre = Alignment(horizontal="center")

    for col_idx, txt in enumerate(entetes, start=1):
        cell = ws.cell(row=3, column=col_idx, value=txt)
        cell.font = Font(bold=True)
        cell.fill = fill
        cell.alignment = centre

    for row_idx, ligne in enumerate(lignes, start=4):
        for col_idx, val in enumerate(ligne, start=1):
            cell = ws.cell(row=row_idx, column=col_idx,
                           value=float(val) if col_idx > 1 else val)
            if col_idx > 1:
                cell.alignment = centre
                cell.number_format = "0.0"

    for col_idx in range(1, len(entetes) + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 14


def _table_pdf(pivot):
    entetes, lignes = _pivot_vers_lignes(pivot)
    data = [entetes] + lignes
    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4472C4")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("ALIGN", (1, 0), (-1, -1), "CENTER"),
        ("ALIGN", (0, 0), (0, -1), "LEFT"),
        ("FONTSIZE", (0, 0), (-1, -1), 10),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1),
         [colors.white, colors.HexColor("#F2F2F2")]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))
    return table


# ============================================================
# Export Word — mono lieu
# ============================================================
def export_word(lieu, coords, nb_annees, pivots, extremes=None) -> io.BytesIO:
    doc = Document()
    titre = doc.add_heading(f"Étude climatique — {lieu}", level=1)
    titre.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(
        f"Coordonnées : {coords['latitude']:.4f}° / {coords['longitude']:.4f}°  "
        f"|  Période : {nb_annees} dernières années  "
        f"|  Source : Open-Meteo (ERA5)"
    )
    run.italic = True
    run.font.size = Pt(10)

    for var_key, pivot in pivots.items():
        config = VARIABLES[var_key]
        doc.add_heading(config["label"], level=2)
        _ajouter_tableau_word(doc, pivot)

    graphique_combo = generer_graphique_ombrothermique(pivots)
    if graphique_combo:
        doc.add_heading("Diagramme ombrothermique", level=2)
        doc.add_picture(graphique_combo, width=Inches(6.0))

    for var_key in ["humidite", "vent", "ensoleillement"]:
        if var_key in pivots:
            config = VARIABLES[var_key]
            doc.add_heading(f"Graphique — {config['label']}", level=2)
            graphique = generer_graphique_simple(pivots[var_key], var_key)
            doc.add_picture(graphique, width=Inches(6.0))

    _ajouter_extremes_word(doc, extremes)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


# ============================================================
# Export Excel — mono lieu
# ============================================================
def export_excel(lieu, coords, nb_annees, pivots, extremes=None) -> io.BytesIO:
    wb = Workbook()
    wb.remove(wb.active)

    for var_key, pivot in pivots.items():
        config = VARIABLES[var_key]
        nom_feuille = config["label"].split(" (")[0][:30]
        ws = wb.create_sheet(nom_feuille)
        _remplir_feuille_excel(ws, f"{config['label']} — {lieu}", pivot)

    _ajouter_extremes_excel(wb, extremes)

    ws_info = wb.create_sheet("Informations")
    infos = [
        ("Lieu", lieu),
        ("Latitude", coords["latitude"]),
        ("Longitude", coords["longitude"]),
        ("Période", f"{nb_annees} dernières années"),
        ("Variables", ", ".join(VARIABLES[v]["label"] for v in pivots)),
        ("Source", "Open-Meteo / ERA5"),
    ]
    for i, (cle, val) in enumerate(infos, start=1):
        ws_info[f"A{i}"] = cle
        ws_info[f"B{i}"] = val
        ws_info[f"A{i}"].font = Font(bold=True)
    ws_info.column_dimensions["A"].width = 18
    ws_info.column_dimensions["B"].width = 50

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


# ============================================================
# Export PDF — mono lieu
# ============================================================
def export_pdf(lieu, coords, nb_annees, pivots, extremes=None) -> io.BytesIO:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4),
                            rightMargin=1.5*cm, leftMargin=1.5*cm,
                            topMargin=1.5*cm, bottomMargin=1.5*cm)

    styles = getSampleStyleSheet()
    style_titre = ParagraphStyle("Titre", parent=styles["Heading1"],
                                  alignment=1, fontSize=18, spaceAfter=12)
    style_sous = ParagraphStyle("Sous", parent=styles["Normal"],
                                 alignment=1, fontSize=10, textColor=colors.grey,
                                 spaceAfter=20)
    style_h2 = ParagraphStyle("H2", parent=styles["Heading2"],
                               fontSize=13, spaceAfter=10)

    elements = []
    elements.append(Paragraph(f"Étude climatique — {lieu}", style_titre))
    elements.append(Paragraph(
        f"Coordonnées : {coords['latitude']:.4f}° / {coords['longitude']:.4f}° "
        f"| Période : {nb_annees} dernières années | Source : Open-Meteo (ERA5)",
        style_sous
    ))

    for var_key, pivot in pivots.items():
        config = VARIABLES[var_key]
        elements.append(Paragraph(config["label"], style_h2))
        elements.append(_table_pdf(pivot))
        elements.append(PageBreak())

    graphique_combo = generer_graphique_ombrothermique(pivots)
    if graphique_combo:
        elements.append(Paragraph("Diagramme ombrothermique", style_h2))
        elements.append(Image(graphique_combo, width=22*cm, height=11*cm))
        elements.append(PageBreak())

    for var_key in ["humidite", "vent", "ensoleillement"]:
        if var_key in pivots:
            config = VARIABLES[var_key]
            elements.append(Paragraph(f"Graphique — {config['label']}", style_h2))
            graphique = generer_graphique_simple(pivots[var_key], var_key)
            elements.append(Image(graphique, width=22*cm, height=9*cm))
            elements.append(PageBreak())

    _ajouter_extremes_pdf(elements, extremes, style_h2)

    doc.build(elements)
    buffer.seek(0)
    return buffer


# ============================================================
# Exports — mode comparaison
# ============================================================
def export_word_comparaison(lieux_noms, nb_annees, pivots_compare,
                              extremes_par_lieu=None) -> io.BytesIO:
    doc = Document()
    titre = doc.add_heading("Étude climatique comparative", level=1)
    titre.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(
        f"Lieux comparés : {', '.join(lieux_noms)}  "
        f"|  Période : {nb_annees} dernières années  "
        f"|  Source : Open-Meteo (ERA5)"
    )
    run.italic = True
    run.font.size = Pt(10)

    for var_key, pivot in pivots_compare.items():
        config = VARIABLES[var_key]
        doc.add_heading(config["label"], level=2)
        _ajouter_tableau_word(doc, pivot)
        graphique = generer_graphique_comparaison(pivot, var_key)
        doc.add_picture(graphique, width=Inches(6.0))

    if extremes_par_lieu:
        for nom_lieu, extremes in extremes_par_lieu.items():
            _ajouter_extremes_word(doc, extremes, f"Extrêmes — {nom_lieu}")

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def export_excel_comparaison(lieux_noms, nb_annees, pivots_compare,
                               extremes_par_lieu=None) -> io.BytesIO:
    wb = Workbook()
    wb.remove(wb.active)

    for var_key, pivot in pivots_compare.items():
        config = VARIABLES[var_key]
        nom_feuille = config["label"].split(" (")[0][:30]
        ws = wb.create_sheet(nom_feuille)
        _remplir_feuille_excel(ws, f"{config['label']} — Comparaison", pivot)

    if extremes_par_lieu:
        for nom_lieu, extremes in extremes_par_lieu.items():
            _ajouter_extremes_excel(wb, extremes, f"Extrêmes {nom_lieu}")

    ws_info = wb.create_sheet("Informations")
    infos = [
        ("Mode", "Comparaison multi-lieux"),
        ("Lieux", ", ".join(lieux_noms)),
        ("Période", f"{nb_annees} dernières années"),
        ("Variables", ", ".join(VARIABLES[v]["label"] for v in pivots_compare)),
        ("Source", "Open-Meteo / ERA5"),
    ]
    for i, (cle, val) in enumerate(infos, start=1):
        ws_info[f"A{i}"] = cle
        ws_info[f"B{i}"] = val
        ws_info[f"A{i}"].font = Font(bold=True)
    ws_info.column_dimensions["A"].width = 18
    ws_info.column_dimensions["B"].width = 60

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def export_pdf_comparaison(lieux_noms, nb_annees, pivots_compare,
                             extremes_par_lieu=None) -> io.BytesIO:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4),
                            rightMargin=1.5*cm, leftMargin=1.5*cm,
                            topMargin=1.5*cm, bottomMargin=1.5*cm)

    styles = getSampleStyleSheet()
    style_titre = ParagraphStyle("Titre", parent=styles["Heading1"],
                                  alignment=1, fontSize=18, spaceAfter=12)
    style_sous = ParagraphStyle("Sous", parent=styles["Normal"],
                                 alignment=1, fontSize=10, textColor=colors.grey,
                                 spaceAfter=20)
    style_h2 = ParagraphStyle("H2", parent=styles["Heading2"],
                               fontSize=13, spaceAfter=10)

    elements = []
    elements.append(Paragraph("Étude climatique comparative", style_titre))
    elements.append(Paragraph(
        f"Lieux comparés : {', '.join(lieux_noms)} | "
        f"Période : {nb_annees} dernières années | Source : Open-Meteo (ERA5)",
        style_sous
    ))

    for var_key, pivot in pivots_compare.items():
        config = VARIABLES[var_key]
        elements.append(Paragraph(config["label"], style_h2))
        elements.append(_table_pdf(pivot))
        elements.append(Spacer(1, 0.5*cm))
        graphique = generer_graphique_comparaison(pivot, var_key)
        elements.append(Image(graphique, width=22*cm, height=10*cm))
        elements.append(PageBreak())

    if extremes_par_lieu:
        for nom_lieu, extremes in extremes_par_lieu.items():
            _ajouter_extremes_pdf(elements, extremes, style_h2,
                                    f"Extrêmes — {nom_lieu}")

    doc.build(elements)
    buffer.seek(0)
    return buffer
